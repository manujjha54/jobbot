"""
profile_builder.py
Turns an uploaded resume (.docx or .pdf) plus a few form fields into a
profile.json the rest of the pipeline (search_jobs, match_engine, ats_score,
apply_approved) already knows how to use. This is what makes the tool work
for ANY resume/role, not just a hardcoded profile.

Nothing here calls out to any AI model - it's the same kind of heuristic
text processing as ats_score.py (frequency counts, section-header detection,
proper-noun spotting), applied to your own resume instead of a job posting.
Always review the suggested skills/titles before saving - it's a starting
point to edit, not a guarantee of accuracy.
"""

import io
import re
from docx import Document
from pypdf import PdfReader

SECTION_HEADER_RE = re.compile(r"^[A-Z][A-Z &/]{3,40}$")
SKILL_SECTION_NAMES = ["SKILL", "COMPETENC", "TOOL", "TECHNOLOG", "EXPERTISE", "PROFICIENC"]
TITLE_WORDS = [
    "manager", "engineer", "director", "lead", "specialist", "analyst",
    "designer", "developer", "consultant", "executive", "officer",
    "coordinator", "representative", "associate", "architect", "scientist",
    "administrator", "strategist", "advisor", "supervisor", "head",
]
YEARS_RE = re.compile(r"(\d{1,2})\+?\s*years?", re.IGNORECASE)

DEFAULT_SENIORITY_KEYWORDS = ["lead", "manager", "head of", "director", "vp", "senior", "principal"]
DEFAULT_EXCLUDE_KEYWORDS = ["unpaid", "internship", "commission only", "no experience required"]

# Generic role/seniority words stripped out when deriving anchor terms from a
# target title, so "Director of Customer Success" becomes the anchor phrase
# "customer success" instead of the whole (much rarer) literal phrase.
GENERIC_TITLE_WORDS = {
    "manager", "director", "lead", "head", "vp", "senior", "principal",
    "of", "the", "team", "officer", "specialist", "associate", "chief",
    "executive", "coordinator", "administrator", "supervisor", "junior",
}

STOPWORDS = set("""
a an the and or but if then else for of to in on at by with without within
is are was were be been being this that these those it its as from into
your you we our their they he she his her them i me my mine ours yours
will would can could should shall may might must have has had do does did
not no yes so than too very just about over under between across per via
role team work working years experience including etc such more most all
any some each other another new job company position looking seeking
""".split())


def extract_text_from_upload(file_bytes, filename):
    """Reads plain text out of a .docx or .pdf resume, given raw bytes and
    the original filename (used only to tell which format it is)."""
    lower = filename.lower()
    stream = io.BytesIO(file_bytes)
    if lower.endswith(".docx"):
        doc = Document(stream)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)
    elif lower.endswith(".pdf"):
        reader = PdfReader(stream)
        parts = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(parts)
    else:
        raise ValueError("Unsupported file type - please upload a .docx or .pdf resume.")


def _find_skill_section_lines(lines):
    """Look for a SKILLS/COMPETENCIES/TOOLS-style header and grab the lines
    directly under it, up to the next all-caps section header."""
    for i, line in enumerate(lines):
        stripped = line.strip()
        if not stripped or not SECTION_HEADER_RE.match(stripped):
            continue
        if any(name in stripped.upper() for name in SKILL_SECTION_NAMES):
            captured = []
            for j in range(i + 1, min(i + 8, len(lines))):
                nxt = lines[j].strip()
                if not nxt:
                    continue
                if SECTION_HEADER_RE.match(nxt) and not any(name in nxt.upper() for name in SKILL_SECTION_NAMES):
                    break  # hit the next section
                captured.append(nxt)
            return captured
    return []


def suggest_skills_from_text(text, max_skills=30):
    lines = text.split("\n")
    section_lines = _find_skill_section_lines(lines)

    skills = set()
    if section_lines:
        raw = " ".join(section_lines)
        # Skills sections are usually comma/bullet/semicolon/pipe separated
        candidates = re.split(r"[,;•|/\n]", raw)
        for c in candidates:
            c = c.strip(" -–:\t")
            c = re.sub(r"\)+$", "", c)  # trailing stray ) from split-on-comma artifacts like "(NRR, GRR)"
            c = re.sub(r"^\(+", "", c)  # leading stray (
            if 2 <= len(c) <= 40 and not c.isdigit():
                skills.add(c)

    if len(skills) < 6:
        # Fallback: frequency-based + proper-noun extraction across the whole resume
        words = re.findall(r"[A-Za-z][A-Za-z+/\-]{2,}", text)
        from collections import Counter
        freq = Counter(w.lower() for w in words if w.lower() not in STOPWORDS and len(w) > 3)
        for word, count in freq.most_common(max_skills):
            if count >= 2:
                skills.add(word)

        titlecase_words = re.findall(r"\b[A-Z][a-zA-Z]{2,}\b", text)
        common_words = {"The", "This", "That", "With", "From", "Have", "Been"}
        for w in titlecase_words:
            if w not in common_words:
                skills.add(w)

    # Dedupe case-insensitively, prefer capitalized form
    best_by_lower = {}
    for s in skills:
        key = s.lower()
        if key not in best_by_lower or (s[0].isupper() and not best_by_lower[key][0].isupper()):
            best_by_lower[key] = s

    result = sorted(best_by_lower.values(), key=lambda s: s.lower())
    return result[:max_skills]


def suggest_titles_from_text(text, max_titles=5):
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    suggestions = []

    def is_section_header(line):
        letters_only = re.sub(r"[^A-Za-z]", "", line)
        return letters_only.isupper() and len(letters_only) > 2

    # Headline: usually one of the first few lines, containing a role word
    for line in lines[:6]:
        if is_section_header(line):
            continue
        if len(line) < 80 and any(w in line.lower() for w in TITLE_WORDS):
            cleaned = re.sub(r"[|•\-–].*$", "", line).strip()
            if cleaned and cleaned not in suggestions:
                suggestions.append(cleaned)

    # Job titles from experience entries: lines containing a role word,
    # reasonably short, not a bullet point (bullets usually start with a verb)
    for line in lines:
        if line in suggestions or is_section_header(line):
            continue
        if len(line) < 70 and any(w in line.lower() for w in TITLE_WORDS) and not line.startswith(("-", "•", "*")):
            if not re.match(r"^(led|managed|drove|built|created|delivered|owned|partner|coach)", line.lower()):
                suggestions.append(line)
        if len(suggestions) >= max_titles:
            break

    # Dedupe while preserving order
    seen = set()
    deduped = []
    for s in suggestions:
        key = s.lower()
        if key not in seen:
            seen.add(key)
            deduped.append(s)
    return deduped[:max_titles]


def suggest_years_experience(text):
    matches = [int(m) for m in YEARS_RE.findall(text)]
    return max(matches) if matches else None


def build_profile(
    resume_filename,
    name,
    email,
    phone,
    linkedin,
    location,
    target_titles,
    must_have_skills,
    nice_to_have_skills,
    years_experience,
    remote_first,
    acceptable_locations,
    open_to_relocation=None,
    timezone_alignment="",
    exclude_keywords=None,
):
    """Assembles a profile.json-compatible dict from form fields."""
    target_titles = [t.strip() for t in target_titles if t.strip()]
    must_have_skills = [s.strip() for s in must_have_skills if s.strip()]
    nice_to_have_skills = [s.strip() for s in nice_to_have_skills if s.strip()]
    acceptable_locations = [l.strip() for l in acceptable_locations if l.strip()]

    # Anchor terms gate out unrelated jobs during matching. Using the full
    # literal title (e.g. "director of customer success") is too strict -
    # that exact phrase rarely appears in a real posting. Instead, strip
    # generic role/seniority words and keep the meaningful core (e.g.
    # "customer success"), which is what actually shows up in job text.
    role_anchor_terms = set()
    for title in target_titles:
        words = re.findall(r"[a-zA-Z]+", title.lower())
        core_words = [w for w in words if w not in GENERIC_TITLE_WORDS]
        if core_words:
            role_anchor_terms.add(" ".join(core_words))
        else:
            role_anchor_terms.add(title.lower())  # fallback if title was all generic words
    role_anchor_terms = list(role_anchor_terms)

    return {
        "name": name,
        "email": email,
        "phone": phone,
        "linkedin": linkedin,
        "location": location,
        "open_to_relocation": open_to_relocation or [],
        "years_experience": years_experience or 0,
        "resume_path": resume_filename,
        "target_titles": target_titles,
        "seniority_keywords": DEFAULT_SENIORITY_KEYWORDS,
        "role_anchor_terms": role_anchor_terms,
        "must_have_any_skills": must_have_skills,
        "nice_to_have_skills": nice_to_have_skills,
        "location_preferences": {
            "remote_first": bool(remote_first),
            "acceptable_locations": acceptable_locations,
            "timezone_alignment": timezone_alignment,
        },
        "exclude_keywords": exclude_keywords if exclude_keywords is not None else DEFAULT_EXCLUDE_KEYWORDS,
    }
